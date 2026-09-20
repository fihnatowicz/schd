"""
Собирает ../index.html из xlsx-расписания.

    python build/build.py                       # берёт первый *.xlsx в папке проекта, лист "1  акт"
    python build/build.py file.xlsx "1  акт"    # явно указать файл и лист

Соседние одинаковые 45-минутные слоты склеиваются в один блок (границы слотов
сохраняются — по ним страница определяет «перерыв» внутри блока).
"""
import glob
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

DAYS = {
    'понедельник': ('Понедельник', 'Пн', 1),
    'вторник': ('Вторник', 'Вт', 2),
    'среда': ('Среда', 'Ср', 3),
    'четверг': ('Четверг', 'Чт', 4),
    'пятница': ('Пятница', 'Пт', 5),
    'суббота': ('Суббота', 'Сб', 6),
}
ROLE = {
    'п': 'проф.', 'д': 'доц.', 'д.': 'доц.',
    'с.пр': 'ст. преп.', 'с.пр.': 'ст. преп.', 'ст.п': 'ст. преп.',
    'пр': 'преп.', 'пр.': 'преп.',
}
KIND = {
    'пр': ('ПР', 'практическое занятие'),
    'инд': ('ИНД', 'индивидуальное занятие'),
    'гр': ('ГР', 'групповое занятие'),
    'л': ('Л', 'лекция'),
    'с': ('С', 'семинар'),
    'л пр': ('Л·ПР', 'лекция / практика'),
    'л/с': ('Л·С', 'лекция / семинар'),
}
SUBJECT_FIX = {'Бел.яз.': 'Белорусский язык'}

META = {
    'title': 'Ангелина',
    'subtitle': 'актёрское искусство · 1 курс',
    'logo': '💅',
    # бегущая лента в шапке
    'ribbon': '💅 АНГЕЛИНА СУПЕР ✨ ANGELINA THE BEST 💖 АНГЕЛИНА THE BEST ⭐ ANGELINA SUPER 🎀 '
              'АНГЕЛИНА ЛУЧШАЯ 💅 ANGELINA SLAY 🌟 АНГЕЛИНА ЗВЕЗДА ♥',
}


def clean(s):
    if s is None:
        return ''
    s = str(s).replace('\xa0', ' ')
    s = re.sub(r'\s+', ' ', s).strip()
    return re.sub(r'[/\s]+$', '', s).strip()


def fix_teacher(t):
    t = clean(t)
    t = re.sub(r',$', '.', t)
    t = re.sub(r'\s*\.\s*', '.', t)
    t = re.sub(r'^(\S+)\.?\s*', r'\1 ', t)
    return t.strip()


def norm_room(s):
    s = clean(s)
    m = re.match(r'^к\.?\s*(\d)\s*,?\s*а\.?\s*(\d+)$', s, re.I)
    if m:
        return {'building': int(m.group(1)), 'room': m.group(2)}
    return {'place': s} if s else None


def norm_kind(k):
    k = clean(k).lower().rstrip('.')
    if not k:
        return None
    if k in KIND:
        code, label = KIND[k]
        return {'code': code, 'label': label}
    return {'code': k.upper(), 'label': k}


def norm_role(r):
    r = clean(r).lower()
    return ROLE.get(r) or ROLE.get(r.rstrip('.')) or (r or None)


def parse_time(t):
    m = re.match(r'(\d{2})\.(\d{2})\s*-\s*(\d{2})\.(\d{2})', clean(t))
    return f'{m.group(1)}:{m.group(2)}', f'{m.group(3)}:{m.group(4)}'


def split_parts(subject, kind, role, teacher, room):
    subject, kind, role, teacher, room = map(clean, (subject, kind, role, teacher, room))
    if '/' in teacher:                       # подгруппы: два преподавателя / две аудитории
        teachers = [fix_teacher(x) for x in teacher.split('/')]
        n = len(teachers)

        def spl(v):
            parts = [clean(x) for x in v.split('/')]
            return parts if len(parts) == n else [v] * n

        subjects, kinds, roles, rooms = spl(subject), spl(kind), spl(role), spl(room)
        return [dict(subject=SUBJECT_FIX.get(subjects[i], subjects[i]), kind=norm_kind(kinds[i]),
                     role=norm_role(roles[i]), teacher=teachers[i], room=norm_room(rooms[i]))
                for i in range(n)]
    return [dict(subject=SUBJECT_FIX.get(subject, subject), kind=norm_kind(kind),
                 role=norm_role(role), teacher=fix_teacher(teacher) or None, room=norm_room(room))]


def special(parts):
    """Особые записи: условные/периодические занятия."""
    s = parts[0]['subject']
    if s.startswith('Единый день информирования'):
        parts[0]['subject'] = 'Единый день информирования'
        return 'только во 2-й вторник месяца', 'second-tuesday'
    if s.startswith('Кураторский час'):
        parts[0]['subject'] = 'Кураторский / информационный час'
        return 'кураторский — 1 ч/мес, информационный — 3 ч/мес', None
    return None, None


def parse(xlsx, sheet):
    ws = openpyxl.load_workbook(xlsx, data_only=True)[sheet]
    days, current = [], None
    for r in range(12, ws.max_row + 1):
        dname = clean(ws.cell(r, 1).value).lower()
        if dname in DAYS:
            full, short, dow = DAYS[dname]
            current = {'dow': dow, 'name': full, 'short': short, 'lessons': []}
            days.append(current)
        if current is None:
            continue
        t = ws.cell(r, 2).value
        if not t or not re.match(r'\d{2}\.\d{2}', clean(t)):
            if days and dname == '' and t is None and r > 84:
                break
            continue
        start, end = parse_time(t)
        if not clean(ws.cell(r, 3).value):
            continue
        parts = split_parts(*(ws.cell(r, c).value for c in range(3, 8)))
        note, cond = special(parts)
        key = json.dumps(parts, ensure_ascii=False, sort_keys=True)
        lessons = current['lessons']
        if lessons and lessons[-1]['_key'] == key:
            lessons[-1]['slots'].append([start, end])
            lessons[-1]['end'] = end
        else:
            lessons.append({'_key': key, 'start': start, 'end': end, 'slots': [[start, end]],
                            'parts': parts, 'note': note, 'cond': cond})
    for d in days:
        for lesson in d['lessons']:
            del lesson['_key']
    return {'meta': META, 'days': days}


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(ROOT, '*.xlsx')))[0]
    sheet = sys.argv[2] if len(sys.argv) > 2 else '1  акт'
    data = parse(xlsx, sheet)
    tpl = open(os.path.join(HERE, 'template.html'), encoding='utf-8').read()
    js = json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    assert tpl.count('__DATA__') == 1
    out = os.path.join(ROOT, 'index.html')      # index.html — чтобы GitHub Pages открывал корень репозитория
    with open(out, 'w', encoding='utf-8', newline='\n') as f:
        f.write(tpl.replace('__DATA__', js))
    total = sum(len(d['lessons']) for d in data['days'])
    print(f'{os.path.basename(xlsx)} [{sheet}] -> {out}: {total} блоков в {len(data["days"])} днях')


if __name__ == '__main__':
    main()
